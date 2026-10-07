from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"F:\Projects\ElectronicCRM")
OUTPUT_DIR = ROOT / "artifacts"
OUTPUT = OUTPUT_DIR / "Инструкция по импорту и обучению df 8832.docx"

IMAGES = {
    "import_results": Path(r"C:\Users\Farianskiy\AppData\Local\Temp\codex-clipboard-68fe3561-8ed9-47f3-9efd-af9533da1bb6.png"),
    "row_editor": Path(r"C:\Users\Farianskiy\AppData\Local\Temp\codex-clipboard-1da78cf0-03ef-4684-ad8d-e5e3aa31370b.png"),
    "learning_workspace": Path(r"C:\Users\Farianskiy\AppData\Local\Temp\codex-clipboard-cf68ffb0-f4e0-4ca5-974a-5d5909ec4fc7.png"),
    "draft_check": Path(r"C:\Users\Farianskiy\AppData\Local\Temp\codex-clipboard-68a67844-4900-4a44-a091-c498e2275b46.png"),
    "literal_fail": Path(r"C:\Users\Farianskiy\AppData\Local\Temp\codex-clipboard-c47ceefd-ae2b-45de-9f51-73e102e6c8b4.png"),
    "numeric_template": Path(r"C:\Users\Farianskiy\AppData\Local\Temp\codex-clipboard-2a5e458b-38d0-4c3a-9b1f-44e5c726e676.png"),
    "match_results": Path(r"C:\Users\Farianskiy\AppData\Local\Temp\codex-clipboard-e87cb990-b25e-4a06-98a9-017d23448260.png"),
}

NAVY = "17365D"
TEAL = "0F766E"
LIGHT_BLUE = "EAF2F8"
PALE_TEAL = "E8F5F3"
LIGHT_GRAY = "F2F4F7"
BORDER = "D9D9D9"
RED = "B42318"
AMBER = "9A6700"
GREEN = "067647"
STEP_COUNTER = 0


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep(paragraph, next_=False) -> None:
    paragraph.paragraph_format.keep_together = True
    if next_:
        paragraph.paragraph_format.keep_with_next = True


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    global STEP_COUNTER
    if level == 1:
        STEP_COUNTER = 0
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.paragraph_format.page_break_before = True
    keep(p, next_=True)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.add_run(text)


def add_number(doc: Document, text: str) -> None:
    global STEP_COUNTER
    STEP_COUNTER += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.paragraph_format.space_after = Pt(2)
    lead = p.add_run(f"{STEP_COUNTER}.  ")
    lead.bold = True
    p.add_run(text)


def add_labeled_paragraph(doc: Document, label: str, text: str, color: str | None = None) -> None:
    p = doc.add_paragraph()
    lead = p.add_run(label + " ")
    lead.bold = True
    if color:
        lead.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.text = text
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8.5)
    for r_index, row_values in enumerate(rows):
        row = table.add_row()
        for i, text in enumerate(row_values):
            cell = row.cells[i]
            cell.text = str(text)
            if r_index % 2:
                set_cell_shading(cell, LIGHT_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_screenshot(doc: Document, key: str, caption: str, width: float = 8.55) -> None:
    path = IMAGES[key]
    if not path.exists():
        p = doc.add_paragraph()
        r = p.add_run(f"Снимок экрана недоступен: {path.name}")
        r.font.color.rgb = RGBColor.from_string(RED)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    picture = run.add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", f"Снимок экрана: {key}")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)
    rr = cap.add_run(caption)
    rr.italic = True
    rr.font.size = Pt(9)
    rr.font.color.rgb = RGBColor(89, 89, 89)


def add_page_break(doc: Document) -> None:
    # Main headings carry page_break_before. Keeping the break on the heading
    # avoids an empty page when the previous section happens to fill a page.
    return


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.04

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        styles[list_name].font.name = "Aptos"
        styles[list_name].font.size = Pt(10.5)


def add_footer(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Electronic CRM   Инструкция по импорту и обучению")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(117, 117, 117)


def build() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(1.1)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    add_footer(section)
    configure_styles(doc)

    # Title page
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("Импорт и обучение распознавания для файла df 8832")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(8)
    sub.paragraph_format.space_after = Pt(16)
    r = sub.add_run("Пошаговая инструкция для Electronic CRM")
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    r.bold = True

    intro = doc.add_paragraph()
    intro.add_run(
        "Инструкция показывает полный путь: загрузить Excel, назначить тип товара, "
        "исправить обязательные характеристики, создать учебные примеры, проверить "
        "правила на всём файле, выпустить версию распознавания и применить импорт к каталогу."
    )
    intro.paragraph_format.space_after = Pt(12)

    add_labeled_paragraph(
        doc,
        "Главное:",
        "обучение не изменяет Excel и не загружает товары автоматически. Оно создаёт правила, "
        "которые CRM применит при повторном анализе строк и в последующих импортах.",
        TEAL,
    )

    add_table(
        doc,
        ["Что находится в файле", "Значение"],
        [
            ["Строк товаров", "95"],
            ["Колонки", "Артикул, наименование, производитель"],
            ["Производитель", "IEK во всех 95 строках"],
            ["Тип товаров", "Силовые автоматические выключатели ВА88"],
            ["Варианты", "3P и 4P, MASTER и обычные, часть строк с электронным расцепителем"],
            ["Отдельные колонки характеристик", "Нет. Значения нужно извлекать из наименования"],
        ],
        [3.2, 6.2],
    )

    add_heading(doc, "Результат после выполнения инструкции", 2)
    add_bullet(doc, "CRM знает, что строки файла относятся к типу «Силовой автомат».")
    add_bullet(doc, "Из названия извлекаются количество полюсов, номинальный ток, ПКС и серия.")
    add_bullet(doc, "Варианты написания обычных товаров и MASTER проверяются отдельно.")
    add_bullet(doc, "После исправления всех обязательных полей пакет можно отправить на проверку и применить к каталогу.")

    add_page_break(doc)
    add_heading(doc, "1 Как CRM понимает этот Excel", 1)
    p = doc.add_paragraph(
        "В файле нет колонок «Количество полюсов», «Номинальный ток», «ПКС» и «Серия». "
        "Поэтому CRM видит исходную строку, но должна извлечь характеристики из текста наименования."
    )
    keep(p)

    add_table(
        doc,
        ["Фрагмент наименования", "Что означает", "Характеристика в CRM", "Что выделять при обучении"],
        [
            ["ВА88-32", "Серия", "Серия товара", "Весь фрагмент ВА88-32"],
            ["3Р", "Три полюса", "Количество полюсов", "Весь фрагмент 3Р"],
            ["100А", "Номинальный ток 100 А", "Номинальный ток", "Только число 100"],
            ["25кА", "Отключающая способность 25 кА", "ПКС", "Только число 25"],
            ["с электрон. расц.", "Электронный расцепитель", "Тип расцепителя", "Фразу электрон. расц."],
            ["MASTER", "Вариант линейки в названии", "Не считать серией без отдельной характеристики", "Обычно не размечать"],
        ],
        [2.0, 2.3, 2.5, 2.8],
    )

    add_heading(doc, "Пример правильного разбора", 2)
    add_labeled_paragraph(doc, "Исходная строка:", "Авт. выкл. ВА88-32 3Р 100А 25кА IEK")
    add_table(
        doc,
        ["Поле", "Значение"],
        [
            ["Производитель", "IEK"],
            ["Вид товара", "Основной товар"],
            ["Тип товара", "Силовой автомат"],
            ["Серия товара", "ВА88-32"],
            ["Количество полюсов", "3Р"],
            ["Номинальный ток", "100 А"],
            ["ПКС", "25 кА"],
            ["Тип расцепителя", "Не заполнять, если в названии он не указан"],
        ],
        [3.0, 6.4],
    )
    add_labeled_paragraph(
        doc,
        "Не додумывайте:",
        "если характеристика отсутствует в названии и в Excel, не назначайте её вручную по предположению.",
        RED,
    )

    add_page_break(doc)
    add_heading(doc, "2 Загрузка файла", 1)
    add_number(doc, "Откройте раздел «Импорт» и создайте новый пакет.")
    add_number(doc, "Выберите режим «Добавить новые товары». Этот режим подходит для df 8832.xlsx, если этих артикулов ещё нет в каталоге.")
    add_number(doc, "Нажмите «Выберите файл», укажите df 8832.xlsx и нажмите «Загрузить и проанализировать».")
    add_number(doc, "На сопоставлении колонок назначьте: «Артикул» → Артикул, «наименование» → Наименование, «производитель» → Производитель.")
    add_number(doc, "Сохраните сопоставление и запустите анализ.")

    add_heading(doc, "Как выбрать режим", 2)
    add_table(
        doc,
        ["Режим", "Когда использовать", "Что произойдёт"],
        [
            ["Добавить новые товары", "Артикулов ещё нет в каталоге", "Создаются новые товары. Совпавший существующий артикул считается ошибкой."],
            ["Обновить характеристики", "Товары уже существуют и нужно дополнить характеристики", "Обновляются только характеристики. Название, цена и остаток не меняются."],
        ],
        [2.2, 3.1, 4.1],
    )
    add_labeled_paragraph(
        doc,
        "Для первого прогона df 8832.xlsx:",
        "используйте «Добавить новые товары». Не переключайтесь на обновление, пока артикулы не появились в рабочем каталоге.",
        TEAL,
    )

    add_page_break(doc)
    add_heading(doc, "3 Почему после анализа показано 95 ошибок", 1)
    add_screenshot(
        doc,
        "import_results",
        "Экран пакета df 8832.xlsx после первичного анализа: производитель определён, тип и обязательные характеристики ещё требуют подтверждения.",
    )
    add_labeled_paragraph(
        doc,
        "Это ожидаемо:",
        "файл не содержит отдельных колонок характеристик, а для типа «Силовой автомат» обязательны как минимум количество полюсов и номинальный ток.",
        AMBER,
    )
    add_bullet(doc, "Ошибка не означает, что файл испорчен.")
    add_bullet(doc, "CRM уже распознала производителя IEK во всех строках.")
    add_bullet(doc, "Сначала назначьте тип товара, затем заполните и разметьте характеристики на нескольких строках.")

    add_page_break(doc)
    add_heading(doc, "4 Исправление первой строки", 1)
    add_screenshot(
        doc,
        "row_editor",
        "Редактор строки 2. Здесь выбирается тип товара и заполняются характеристики, которые CRM не смогла определить автоматически.",
    )
    add_number(doc, "У строки 2 нажмите «Редактировать».")
    add_number(doc, "В блоке «Тип товара» оставьте «Основные товары» и выберите «Силовой автомат · POWER_CIRCUIT_BREAKER».")
    add_number(doc, "Заполните значения: количество полюсов 3Р, номинальный ток 100, ПКС 25, серия ВА88-32.")
    add_number(doc, "Не заполняйте тип расцепителя, потому что в этой строке он не указан.")
    add_labeled_paragraph(
        doc,
        "Важно:",
        "простое заполнение полей исправит только эту строку. Чтобы CRM научилась распознавать следующие строки, нужно дополнительно связать значение с фрагментом наименования.",
        TEAL,
    )

    add_page_break(doc)
    add_heading(doc, "5 Создание учебного примера", 1)
    p = doc.add_paragraph(
        "Для каждой характеристики выполняется одна и та же последовательность. Ниже приведён пример для номинального тока 100 А."
    )
    add_number(doc, "В редакторе строки прокрутите до блока «Разметка учебного примера».")
    add_number(doc, "Выберите характеристику «Номинальный ток».")
    add_number(doc, "В тексте наименования выделите мышью только цифры 100. Букву А не выделяйте.")
    add_number(doc, "Нажмите «Связать выделенный фрагмент».")
    add_number(doc, "Повторите разметку для остальных характеристик этой строки.")
    add_number(doc, "Нажмите «Сохранить строку». Это сохраняет исправленные значения и выделенные фрагменты.")
    add_number(doc, "После сохранения нажмите «Подтвердить для обучения» отдельно для каждой размеченной характеристики.")

    add_table(
        doc,
        ["Характеристика", "Значение в форме", "Что выделить в названии"],
        [
            ["Количество полюсов", "3Р", "3Р"],
            ["Номинальный ток", "100", "100"],
            ["ПКС", "25", "25"],
            ["Серия товара", "ВА88-32", "ВА88-32"],
        ],
        [2.6, 2.4, 4.4],
    )
    add_labeled_paragraph(
        doc,
        "Порядок кнопок:",
        "сначала «Связать выделенный фрагмент», затем «Сохранить строку», затем «Подтвердить для обучения». Если строка не сохранена, подтверждение будет отклонено.",
        TEAL,
    )

    add_page_break(doc)
    add_heading(doc, "6 Сколько примеров нужно создать", 1)
    p = doc.add_paragraph(
        "Для числового шаблона CRM требует минимум два подтверждённых примера с разными значениями и одинаковым окружением числа. "
        "Один пример показывает правильный ответ, но не доказывает, какая часть названия меняется."
    )

    add_heading(doc, "Минимальный набор для первого рабочего правила", 2)
    add_table(
        doc,
        ["Группа названий", "Пример 1", "Пример 2", "Зачем"],
        [
            ["Обычные 3Р", "Строка 2: 100А", "Строка 4: 12,5А или строка 5: 125А", "Правило для «Авт. выкл. ... 3Р ... IEK»"],
            ["MASTER 3Р", "Строка 3: 100А", "Строка 6: 125А", "Отдельный порядок слов и слово MASTER"],
            ["Обычные 4Р", "Любые две строки 4Р с разным током", "Вторая строка того же формата", "Шаблон 3Р не покрывает 4Р"],
            ["Электронный расцепитель", "Две строки с «электрон. расц.»", "Разные токи при одинаковом формате", "Отдельное окончание названия"],
        ],
        [2.1, 2.4, 2.5, 3.0],
    )
    add_labeled_paragraph(
        doc,
        "Лучше начать:",
        "с двух целых токов, например 100 и 125. Значение 12,5 содержит дробь и не подходит для текущего генератора целочисленных шаблонов.",
        AMBER,
    )
    add_heading(doc, "Почему MASTER обучается отдельно", 2)
    doc.add_paragraph(
        "CRM сравнивает не только число, но и окружение. Названия «Авт. выкл. ... IEK» и «Выкл. авт. ... MASTER IEK» имеют другой порядок слов и дополнительное слово MASTER. "
        "Поэтому правило, созданное для обычной строки, не обязано работать на MASTER."
    )

    add_page_break(doc)
    add_heading(doc, "7 Переход в раздел обучения", 1)
    add_screenshot(
        doc,
        "learning_workspace",
        "Раздел «Распознавание». Выбраны производитель IEK, тип «Силовой автомат» и пакет df 8832.xlsx.",
    )
    add_number(doc, "Откройте «Настройка и качество» → «Распознавание».")
    add_number(doc, "Выберите производителя IEK.")
    add_number(doc, "Выберите тип товара «Силовой автомат».")
    add_number(doc, "Откройте вкладку «Правила» и этап «1. Подготовить черновики».")
    add_number(doc, "В блоке «Мои пакеты импорта» выберите df 8832.xlsx.")
    add_number(doc, "Выберите характеристику, с которой работаете. Для начала — «Номинальный ток».")

    add_page_break(doc)
    add_heading(doc, "8 Проверка подготовленных примеров", 1)
    add_screenshot(
        doc,
        "draft_check",
        "Результат проверки подтверждений. Здесь видно, какие строки поддерживают правило и какие фрагменты не найдены.",
    )
    add_labeled_paragraph(
        doc,
        "Фраза «Состав подтверждений не изменился» означает:",
        "кнопка перепроверила уже сохранённые примеры, но новых подтверждений после последнего запуска не появилось.",
        TEAL,
    )
    add_bullet(doc, "«Фрагмент не найден» — выбранный шаблон не смог повторно найти размеченное значение в названии.")
    add_bullet(doc, "«Предложенное значение: не предлагается» — текущего правила для этой строки ещё нет.")
    add_bullet(doc, "Проверка ничего не сохраняет и не активирует. Это только диагностика.")

    add_page_break(doc)
    add_heading(doc, "9 Почему точное правило для числа не проходит", 1)
    add_screenshot(
        doc,
        "literal_fail",
        "Точные предложения 100 → 100, 16 → 16 и 40 → 40 не проходят, потому что число находится внутри обозначения 100А, 16А или 40А.",
    )
    doc.add_paragraph(
        "Это не обязательно означает неправильную разметку. Точное текстовое правило ищет отдельный самостоятельный фрагмент. "
        "В названии ток записан слитно с единицей измерения, поэтому для него нужен структурный числовой шаблон."
    )
    add_number(doc, "Не сохраняйте точное правило 100 → 100 для номинального тока.")
    add_number(doc, "Прокрутите ниже до блока «Структурные предложения: целое число».")
    add_number(doc, "Нажмите «Показать структурные предложения».")

    add_page_break(doc)
    add_heading(doc, "10 Создание числового шаблона", 1)
    add_screenshot(
        doc,
        "numeric_template",
        "Структурное предложение для тока: CRM заменяет подтверждённое число переменной и проверяет шаблон на учебных примерах.",
    )
    add_number(doc, "Найдите предложение, у которого есть минимум два разных подтверждённых значения и нет конфликтов.")
    add_number(doc, "Нажмите «Проверить на этом импорте».")
    add_number(doc, "Просмотрите строки, где предложено значение, и строки «Название не подходит под шаблон».")
    add_number(doc, "Если шаблон извлекает правильные числа только из нужной группы названий, нажмите «Сохранить числовой шаблон».")
    add_labeled_paragraph(
        doc,
        "Сохранение шаблона:",
        "создаёт черновик. Оно ещё не включает правило в рабочее распознавание.",
        TEAL,
    )

    add_page_break(doc)
    add_heading(doc, "11 Как читать результат проверки", 1)
    add_screenshot(
        doc,
        "match_results",
        "Обычные названия подходят под созданный шаблон, а варианты MASTER не подходят из-за другого порядка слов и дополнительного слова.",
        width=7.55,
    )
    add_table(
        doc,
        ["Сообщение", "Что оно означает", "Что делать"],
        [
            ["Предложено значение 32", "Шаблон нашёл 32А и извлёк число 32", "Проверить правильность и оставить"],
            ["Совпадает с текущим значением", "Шаблон дал уже сохранённое правильное значение", "Хороший результат"],
            ["Название не подходит под шаблон", "Структура названия отличается", "Создать отдельные примеры и отдельный шаблон для этой группы"],
            ["Не предлагается", "Правило не сработало", "Не применять значение вручную без проверки"],
        ],
        [2.8, 4.0, 3.4],
    )
    add_labeled_paragraph(
        doc,
        "В показанном примере:",
        "обычные строки «Авт. выкл. ...» распознаются, а строки «Выкл. авт. ... MASTER ...» требуют отдельного шаблона.",
        AMBER,
    )

    add_page_break(doc)
    add_heading(doc, "12 Составление версии и включение правила", 1)
    add_number(doc, "После сохранения черновиков откройте этап «2. Составить версию».")
    add_number(doc, "Введите понятное название, например «IEK Силовые автоматы df8832 v1».")
    add_number(doc, "Выберите тип черновика «Одиночные числовые шаблоны» и характеристику «Номинальный ток».")
    add_number(doc, "Нажмите «Обновить доступные черновики», отметьте проверенные шаблоны и создайте версию.")
    add_number(doc, "Откройте этап «3. Версии и выпуск» и запустите оценку версии на df 8832.xlsx.")
    add_number(doc, "Проверьте регрессии и конфликты. Если неверные значения не появились, подтвердите выпуск.")
    add_number(doc, "Нажмите «Активировать эту версию».")
    add_labeled_paragraph(
        doc,
        "Только после активации:",
        "правило начинает участвовать в новом анализе. Сохранённый черновик сам по себе ничего не распознаёт.",
        GREEN,
    )
    add_heading(doc, "Повторите для остальных характеристик", 2)
    add_table(
        doc,
        ["Характеристика", "Рекомендуемый подход"],
        [
            ["Количество полюсов", "Подтвердить 3Р и 4Р на нескольких разных строках. Проверить варианты кириллица Р и латиница P."],
            ["ПКС", "Создать числовые примеры 25, 35, 50 и 70; выделять только число перед кА."],
            ["Серия товара", "Создать текстовые примеры ВА88-32, ВА88-33, ВА88-35 и других серий."],
            ["Тип расцепителя", "Обучать только на строках, где прямо написано «электрон. расц.»; остальные не заполнять по догадке."],
        ],
        [3.0, 7.2],
    )

    add_page_break(doc)
    add_heading(doc, "13 Возврат к импорту и применение в каталог", 1)
    add_number(doc, "Вернитесь в пакет df 8832.xlsx.")
    add_number(doc, "Запустите повторный анализ или обновление подсветки, чтобы активная версия распознавания применилась к строкам.")
    add_number(doc, "Откройте фильтр «Ошибки» и проверьте оставшиеся строки.")
    add_number(doc, "Строки с необычным названием исправьте вручную или добавьте новые учебные примеры.")
    add_number(doc, "Добейтесь значения «С ошибками: 0».")
    add_number(doc, "Нажмите «Отправить на проверку».")
    add_number(doc, "В очереди проверки откройте пакет, выберите «Применить в каталог» и подтвердите действие.")
    add_labeled_paragraph(
        doc,
        "Перед применением:",
        "проверьте артикулы. В режиме добавления новых существующий артикул остановит применение всего пакета.",
        RED,
    )

    add_heading(doc, "Проверка результата", 2)
    add_bullet(doc, "Откройте каталог и найдите артикул SVA10-3-0100.")
    add_bullet(doc, "Убедитесь, что тип — «Силовой автомат».")
    add_bullet(doc, "Проверьте: 3Р, 100 А, 25 кА и ВА88-32.")
    add_bullet(doc, "Проверьте одну строку MASTER и одну строку 4Р.")
    add_bullet(doc, "Проверьте строку с электронным расцепителем отдельно.")

    add_page_break(doc)
    add_heading(doc, "14 Практический план именно для df 8832.xlsx", 1)
    add_table(
        doc,
        ["Этап", "Что сделать", "Готово, когда"],
        [
            ["1", "Загрузить файл в режиме добавления новых", "Создан пакет на 95 строк"],
            ["2", "Назначить IEK и тип «Силовой автомат»", "У строк указан правильный тип"],
            ["3", "Разметить две обычные 3Р строки с разными токами", "Есть подтверждения для тока, полюсов, ПКС и серии"],
            ["4", "Создать и проверить структурный шаблон", "Обычные строки получают правильный ток"],
            ["5", "Разметить две MASTER 3Р строки", "MASTER распознаётся отдельным шаблоном"],
            ["6", "Повторить для 4Р и электронного расцепителя", "Основные форматы файла покрыты"],
            ["7", "Собрать, оценить и активировать версию", "Версия имеет статус активной"],
            ["8", "Перезапустить анализ пакета", "Количество ошибок заметно сократилось"],
            ["9", "Исправить исключения вручную", "Ошибок 0"],
            ["10", "Отправить на проверку и применить", "Товары появились в каталоге"],
        ],
        [0.7, 6.2, 3.3],
    )

    add_heading(doc, "Что текущая система не делает автоматически", 2)
    add_bullet(doc, "Не обучается одной кнопкой на всех 95 строках без разметки.")
    add_bullet(doc, "Не считает обычные и MASTER названия одним шаблоном, если структура текста отличается.")
    add_bullet(doc, "Не должна придумывать отсутствующий тип расцепителя.")
    add_bullet(doc, "Целочисленный шаблон не покрывает значение 12,5 без отдельной поддержки дробей.")
    add_bullet(doc, "Черновик не работает до составления, проверки и активации версии.")

    add_page_break(doc)
    add_heading(doc, "15 Краткая памятка по ошибкам", 1)
    add_table(
        doc,
        ["Ситуация", "Причина", "Решение"],
        [
            ["95 ошибок сразу после загрузки", "Нет характеристик в отдельных колонках", "Назначить тип и начать разметку примеров"],
            ["Точное правило не прошло", "Число находится внутри 100А или 25кА", "Использовать структурное числовое предложение"],
            ["Недостаточно разных подтверждённых чисел", "Есть только одно значение", "Подтвердить вторую строку того же формата с другим числом"],
            ["MASTER не распознаётся", "Другой порядок слов и дополнительное слово", "Сделать отдельный шаблон для MASTER"],
            ["3Р работает, 4Р нет", "Шаблон запомнил 3Р в окружении", "Создать отдельные примеры для 4Р"],
            ["После сохранения ничего не изменилось", "Сохранён только черновик", "Составить версию, оценить и активировать"],
            ["Строка всё ещё с ошибкой", "Не заполнена обязательная характеристика", "Открыть редактор и проверить красные обязательные поля"],
            ["Значение 12,5 не входит в шаблон", "Генератор работает с целыми числами", "Оставить ручную обработку или доработать распознавание дробей"],
        ],
        [3.0, 3.5, 3.7],
    )

    add_heading(doc, "Итог", 1)
    doc.add_paragraph(
        "Для df 8832.xlsx сначала достаточно обучить несколько устойчивых групп названий, а не исправлять все 95 строк по одной. "
        "Но обучение нужно выполнять отдельно для разных структур: обычные, MASTER, 3Р, 4Р и строки с электронным расцепителем. "
        "После активации правил повторный анализ заполнит совпадающие строки; оставшиеся исключения исправляются вручную или становятся новыми учебными примерами."
    )

    doc.core_properties.title = "Импорт и обучение распознавания для файла df 8832"
    doc.core_properties.subject = "Electronic CRM"
    doc.core_properties.author = "Electronic CRM"
    doc.core_properties.keywords = "импорт, обучение, распознавание, Excel, IEK, ВА88"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
